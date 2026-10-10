"""The supplier portal (PF8): sign-up and verification, the catalogue,
geometry checks, the price grid, spreadsheet imports and the template, the
inbox, analytics, the listing plan through PF2's interface, the team and
supplier keys, and every route's auth."""

import io
import sys
import types
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app import mail
from app.database import SessionLocal
from app.market import commissions, importer, seed
from app.market.models import (
    Commission,
    FeedRun,
    Product,
    VariantAvailability,
    VariantPrice,
)
from app.models import Payment, Subscription, User
from app.supplier import applications, geometry_check, imports, listing
from app.supplier.models import ListingSubscription, MarketEventDaily, SupplierImport

from .conftest import signup
from .licence_helpers import Clock, error_of
from .market_helpers import (
    CONTRACT,
    FIXTURES,
    SOFA,
    SUPPLIER_ID,
    FakeGateway,
    order_body,
    payments_on,
    png,
    stripe_post,
)
from .supplier_helpers import (
    FEED_CSV,
    admin_headers,
    application,
    apply,
    as_supplier,
    glb,
    no_internet,
    outbox,
    pdf,
    rows_of,
    supplier_key,
    verified,
)


@pytest.fixture(autouse=True)
def _clear_outbox():
    mail.OUTBOX.clear()


def _join(client, owner: dict, email: str, role: str) -> dict:
    """Invite `email` with `role` and accept as a new account; its headers."""
    res = client.post("/supplier/members", json={"email": email, "role": role}, headers=owner)
    assert res.status_code == 201, res.text
    token = outbox("invited you", to=email)[-1].text.split("token=")[1].split()[0]
    h = signup(client, email=email)
    assert client.post("/supplier/invites/accept", json={"token": token}, headers=h).status_code == 200
    return h


def _fixture_supplier(client, email: str = "owner@nord.example.com", **kw) -> dict:
    """The contract fixtures' supplier (verified, six products) with a portal owner."""
    h = signup(client, email=email)
    with SessionLocal() as db:
        owner = db.scalar(select(User).where(User.email == email))
        seed.load_fixture(db, FIXTURES, owner=owner, **kw)
    return as_supplier(h, SUPPLIER_ID)


# --- Sign-up and verification -------------------------------------------------------------------


def test_supplier_application_flow(client, monkeypatch):
    # Anonymous: 401.
    error_of(client.post("/supplier/applications", json=application()), 401, "unauthenticated")
    h = signup(client, email="owner@nord.example.com")
    assert client.get("/supplier/me", headers=h).json()["memberships"] == []
    # Field rules: 422.
    for bad in (
        {"country": "GBR"},
        {"website": "http://nord.example.com"},
        {"regions": []},
        {"regions": ["XX"]},
        {"contact": {"name": "A", "email": "not-an-email"}},
        {"legal_name": ""},
    ):
        error_of(client.post("/supplier/applications", json={**application(), **bad}, headers=h), 422, "validation_failed")
    res = client.post("/supplier/applications", json=application(), headers=h)
    assert res.status_code == 201, res.text
    app = res.json()
    assert app["state"] == "applied" and app["fields"]["regions"] == ["AE", "GB"]
    sid = app["supplier_id"]
    # The "application received" e-mail, to the contact and the applicant.
    received = outbox("We received your supplier application")
    assert {m.to for m in received} == {"orders@nord.example.com", "owner@nord.example.com"}
    assert "Nord Living" in received[0].text and "/supplier/" in received[0].text
    # One application per account: 409.
    error_of(client.post("/supplier/applications", json=application(), headers=h), 409, "already_applied")
    me = client.get("/supplier/me", headers=h).json()
    assert me["role"] == "owner" and me["supplier"]["status"] == "applied" and me["supplier"]["regions"] == ["GB", "AE"]

    # The company document: a PDF ≤ 10 MB.
    owner = as_supplier(h, sid)
    res = client.post("/supplier/documents", files={"file": ("registration.pdf", pdf(), "application/pdf")}, headers=owner)
    assert res.status_code == 201, res.text
    error_of(client.post("/supplier/documents", files={"file": ("x.pdf", b"PK\x03\x04zip", "application/pdf")}, headers=owner), 422, "not_pdf")
    monkeypatch.setattr(applications, "MAX_DOCUMENT_BYTES", 1000)
    error_of(client.post("/supplier/documents", files={"file": ("big.pdf", pdf(5000), "application/pdf")}, headers=owner), 413, "too_large")
    monkeypatch.setattr(applications, "MAX_DOCUMENT_BYTES", 10 * 1024 * 1024)
    assert client.get("/supplier/me", headers=owner).json()["application"]["documents"][0]["name"] == "registration.pdf"

    # Unverified: products can be prepared but not submitted.
    pid = client.post("/supplier/products", json={"sku": "S1", "name": "Stool", "category": "furniture/seating"}, headers=owner).json()["product_id"]
    error_of(client.post(f"/supplier/products/{pid}/submit", headers=owner), 409, "not_verified")

    # The admin reads the application (documents by signed link) and verifies it.
    error_of(client.get(f"/admin/market/suppliers/{sid}/application", headers=owner), 403, "forbidden")
    admin = admin_headers(client)
    seen = client.get(f"/admin/market/suppliers/{sid}/application", headers=admin).json()
    assert seen["state"] == "applied" and seen["fields"]["company_number"] == "01234567"
    doc = client.get(seen["documents"][0]["url"].replace("https://api.truebex.com", ""))
    assert doc.status_code == 200 and doc.content.startswith(b"%PDF-")
    assert client.patch(f"/admin/market/suppliers/{sid}", json={"status": "verified"}, headers=admin).status_code == 200
    assert outbox("Nord Living is verified on Truebex")
    me = client.get("/supplier/me", headers=owner).json()
    assert me["supplier"]["status"] == "verified" and me["application"]["state"] == "verified"

    # Another applicant is declined with a reason (suspending an application).
    h2, sid2 = apply(client, email="second@example.com", name="Dune Interiors")
    res = client.patch(f"/admin/market/suppliers/{sid2}", json={"status": "suspended", "reason": "No company document."}, headers=admin)
    assert res.status_code == 200
    declined = outbox("Your supplier application for Dune Interiors")
    assert declined and "No company document." in declined[0].text
    assert client.get("/supplier/me", headers=h2).json()["application"]["state"] == "declined"
    # A suspended supplier cannot change its catalogue.
    error_of(client.post("/supplier/products", json={"sku": "S1", "name": "Stool", "category": "furniture/seating"}, headers=h2), 403, "supplier_suspended")


# --- Catalogue ---------------------------------------------------------------------------------


def test_supplier_products_crud_and_review(client):
    owner, sid = verified(client)
    # Create (a draft), validation, a variant, a picture, geometry and a price.
    error_of(client.post("/supplier/products", json={"sku": "A", "name": "A", "category": "furniture/hovercraft"}, headers=owner), 422, "unknown_category")
    error_of(client.post("/supplier/products", json={"sku": "", "name": "A", "category": "furniture/seating"}, headers=owner), 422, "validation_failed")
    res = client.post(
        "/supplier/products",
        json={"sku": "CHAIR-1", "name": "Bergen chair", "category": "furniture/seating/armchairs", "description": "Beech."},
        headers=owner,
    )
    assert res.status_code == 201, res.text
    product = res.json()
    pid = product["product_id"]
    assert product["status"] == "draft" and product["editable"]
    error_of(client.post("/supplier/products", json={"sku": "CHAIR-1", "name": "x", "category": "furniture/seating"}, headers=owner), 409, "sku_taken")
    # Incomplete: 422 with what is missing.
    body = error_of(client.post(f"/supplier/products/{pid}/submit", headers=owner), 422, "incomplete")
    assert set(body["data"]["missing"]) >= {"variants", "images", "prices"}
    res = client.post(
        f"/supplier/products/{pid}/variants",
        json={"variant_id": "olive", "options": {"colour": "Olive"}, "dims_mm": [780, 760, 820], "materials": ["beech"], "gtin": "9501101530027"},
        headers=owner,
    )
    assert res.status_code == 201, res.text
    error_of(client.post(f"/supplier/products/{pid}/variants", json={"variant_id": "olive"}, headers=owner), 409, "variant_taken")
    error_of(client.post(f"/supplier/products/{pid}/variants", json={"variant_id": "x", "gtin": "9501101530028"}, headers=owner), 422, "validation_failed")
    res = client.post(f"/supplier/products/{pid}/variants/olive/images", files={"image": ("a.png", png(), "image/png")}, headers=owner)
    assert res.status_code == 201, res.text
    assert res.json()["thumbnail_url"].endswith(".jpg")
    error_of(client.post(f"/supplier/products/{pid}/variants/olive/images", files={"image": ("s.png", png(size=(300, 300)), "image/png")}, headers=owner), 422, "image_too_small")
    res = client.post("/supplier/geometry", data={"product_id": pid, "variant_id": "olive"}, files={"file": ("chair.glb", glb(), "model/gltf-binary")}, headers=owner)
    assert res.status_code == 200, res.text
    assert res.json()["geometry"]["format"] == "glb" and res.json()["check"]["warnings"] == []
    res = client.put("/supplier/prices", params={"region": "GB"}, json={"rows": [{"sku": "CHAIR-1", "variant_id": "olive", "price": "649.00", "stock": "7"}]}, headers=owner)
    assert res.status_code == 200, res.text
    # Submit → in review; a draft cannot be hidden.
    res = client.post(f"/supplier/products/{pid}/submit", headers=owner)
    assert res.status_code == 200 and res.json()["status"] == "pending_review"
    error_of(client.patch(f"/supplier/products/{pid}", json={"name": "x"}, headers=owner), 409, "not_editable")
    # A viewer reads but cannot change the catalogue.
    viewer = as_supplier(_join(client, owner, "viewer@example.com", "viewer"), sid)
    assert client.get(f"/supplier/products/{pid}", headers=viewer).status_code == 200
    body = error_of(client.post("/supplier/products", json={"sku": "V", "name": "V", "category": "furniture/seating"}, headers=viewer), 403, "forbidden")
    assert body["data"]["role"] == "viewer"
    error_of(client.post(f"/supplier/products/{pid}/hide", headers=viewer), 403, "forbidden")
    # Another supplier's product: 404.
    other, _ = verified(client, email="other@example.com", name="Other Co")
    error_of(client.get(f"/supplier/products/{pid}", headers=other), 404, "not_found")
    error_of(client.patch(f"/supplier/products/{pid}", json={"name": "x"}, headers=other), 404, "not_found")

    # Approved: a description change stays approved; a name change goes back to review.
    admin = admin_headers(client)
    assert client.post(f"/admin/market/products/{pid}/approve", headers=admin).status_code == 200
    assert outbox("Approved: Bergen chair")
    res = client.patch(f"/supplier/products/{pid}", json={"description": "Solid beech."}, headers=owner)
    assert res.json()["status"] == "approved"
    res = client.patch(f"/supplier/products/{pid}", json={"name": "Bergen lounge chair"}, headers=owner)
    assert res.json()["status"] == "pending_review"
    client.post(f"/admin/market/products/{pid}/reject", json={"reason": "Picture too dark."}, headers=admin)
    assert "Picture too dark." in outbox("Changes needed: Bergen lounge chair")[0].text
    assert client.post(f"/supplier/products/{pid}/submit", headers=owner).json()["status"] == "pending_review"
    client.post(f"/admin/market/products/{pid}/approve", headers=admin)
    # Hide, show, discontinue, withdraw.
    assert client.post(f"/supplier/products/{pid}/hide", headers=owner).json()["status"] == "hidden"
    error_of(client.get(f"/market/products/{pid}", headers=CONTRACT), 410, "product_withdrawn")
    assert client.post(f"/supplier/products/{pid}/show", headers=owner).json()["status"] == "approved"
    assert client.get(f"/market/products/{pid}", headers=CONTRACT).status_code == 200
    res = client.post(f"/supplier/products/{pid}/discontinue", headers=owner)
    assert res.json()["variants"][0]["status"] == "discontinued"
    with SessionLocal() as db:
        assert db.scalar(select(VariantAvailability.state).where(VariantAvailability.product_id == pid)) == "discontinued"
    assert client.post(f"/supplier/products/{pid}/withdraw", headers=owner).json()["status"] == "withdrawn"
    error_of(client.post(f"/supplier/products/{pid}/nonsense", headers=owner), 404, "not_found")
    listed = client.get("/supplier/products", params={"status": "withdrawn"}, headers=owner).json()
    assert [p["sku"] for p in listed["products"]] == ["CHAIR-1"] and listed["counts"] == {"withdrawn": 1}


def test_supplier_geometry_checks(client, monkeypatch):
    owner, _sid = verified(client)
    pid = client.post("/supplier/products", json={"sku": "C", "name": "Chair", "category": "furniture/seating"}, headers=owner).json()["product_id"]
    client.post(f"/supplier/products/{pid}/variants", json={"variant_id": "default", "dims_mm": [780, 760, 820]}, headers=owner)

    def up(data: bytes, name: str):
        return client.post("/supplier/geometry", data={"product_id": pid, "variant_id": "default"}, files={"file": (name, data)}, headers=owner)

    # Over 500 000 triangles: 422.
    body = error_of(up(glb(triangles=500_001), "chair.glb"), 422, "too_many_triangles")
    assert body["data"]["triangles"] == 500_001
    assert up(glb(triangles=500_000), "chair.glb").status_code == 200
    # The bounding box 10 % off dims_mm: a warning, still stored.
    res = up(glb(size_m=(0.78 * 1.1, 0.76, 0.82)), "chair.glb")
    assert res.status_code == 200
    warning = res.json()["check"]["warnings"][0]
    assert warning["code"] == "size_differs" and warning["measured_mm"] == [858, 760, 820]
    # Within 5 %: no warning (node scale is applied).
    res = up(glb(size_m=(0.39, 0.38, 0.41), nodes_scale=[2, 2, 2]), "chair.glb")
    assert res.json()["check"]["warnings"] == [] and res.json()["check"]["measured_mm"] == [780, 760, 820]
    assert res.json()["geometry"]["rev"] >= 2
    # A renamed ZIP, a name that disagrees with the bytes, compressed meshes: 422.
    error_of(up(b"PK\x03\x04" + b"\x00" * 200, "chair.glb"), 422, "bad_geometry_format")
    error_of(up(b"v 0 0 0\nv 1 0 0\nv 0 1 0\nf 1 2 3\n", "chair.glb"), 422, "bad_geometry_format")
    error_of(up(b"not a model at all", "chair.fbx"), 422, "bad_geometry_format")
    error_of(up(glb(extensions_required=["KHR_draco_mesh_compression"]), "chair.glb"), 422, "unsupported_extension")
    # OBJ by its bytes; a Truebex object (.tbxa).
    res = up(b"# chair\nv 0 0 0\nv 1 0 0\nv 0 1 0\nv 1 1 0\nf 1 2 3 4\n", "chair.obj")
    assert res.status_code == 200 and res.json()["check"]["triangles"] == 2
    res = up((FIXTURES / "geometry" / "sofa-oslo.tbxa").read_bytes(), "sofa.tbxa")
    assert res.status_code == 200 and res.json()["geometry"]["format"] == "tbxa"
    # Over 100 MB: 413.
    monkeypatch.setattr(geometry_check, "MAX_BYTES", 1000)
    error_of(up(glb() + b" " * 2000, "chair.glb"), 413, "too_large")


# --- Prices ----------------------------------------------------------------------------------------


def test_supplier_price_grid_validation(client):
    owner = _fixture_supplier(client)
    grid = client.get("/supplier/prices", params={"region": "GB"}, headers=owner).json()
    assert grid["region"]["currency"] == "GBP" and set(grid["served"]) == {"GB", "AE"}
    oat = next(r for r in grid["rows"] if r["sku"] == "SOFA-OSLO-3" and r["variant_id"] == "oat-linen")
    assert oat["price"] == "1299.00" and oat["amount"] == 129900

    def put(region, rows):
        return client.put("/supplier/prices", params={"region": region}, json={"rows": rows}, headers=owner)

    # Three decimals in GBP → bad_price; USD in GB → currency_mismatch; nothing saved.
    body = error_of(
        put("GB", [
            {"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1299.999"},
            {"sku": "SOFA-OSLO-3", "variant_id": "charcoal-wool", "price": "1449.00", "currency": "USD"},
            {"sku": "SOFA-OSLO-3", "variant_id": "nope", "price": "1.00"},
            {"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1.00", "stock": "lots"},
        ]),
        422,
        "rows_invalid",
    )  # fmt: skip
    codes = {(e["row"], e["column"], e["code"]) for e in body["data"]["errors"]}
    assert {(1, "price", "bad_price"), (2, "currency", "currency_mismatch"), (3, "variant_id", "unknown_variant"), (4, "stock", "bad_integer"), (4, "variant_id", "duplicate_row")} <= codes
    assert client.get("/supplier/prices", params={"region": "GB"}, headers=owner).json()["rows"] == grid["rows"]
    # AED has two decimals too; a thousands separator is refused.
    error_of(put("AE", [{"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1,299.00"}]), 422, "rows_invalid")
    # A region the supplier does not sell in.
    error_of(client.get("/supplier/prices", params={"region": "US"}, headers=owner), 422, "validation_failed")
    # Bulk edit: two variants, price and stock; stock 0 with a lead time → made to order.
    res = put("GB", [
        {"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1249.00", "stock": 3},
        {"sku": "SOFA-OSLO-3", "variant_id": "charcoal-wool", "price": "1449.00", "stock": "0", "lead_time_days": "42"},
    ])  # fmt: skip
    assert res.status_code == 200, res.text
    rows = {r["variant_id"]: r for r in res.json()["rows"] if r["sku"] == "SOFA-OSLO-3"}
    assert (rows["oat-linen"]["price"], rows["oat-linen"]["stock"], rows["oat-linen"]["availability"]) == ("1249.00", 3, "in_stock")
    assert (rows["charcoal-wool"]["availability"], rows["charcoal-wool"]["lead_time_days"]) == ("made_to_order", 42)
    with SessionLocal() as db:
        p = db.scalar(select(VariantPrice).where(VariantPrice.product_id == SOFA["product_id"], VariantPrice.variant_id == "oat-linen", VariantPrice.region == "GB"))
        assert (p.amount, p.source, p.includes_tax, p.tax_rate_bp) == (124900, "manual", True, 2000)
    # The app sees the new price (5.4).
    body = client.get(f"/market/products/{SOFA['product_id']}", params={"region": "GB"}, headers=CONTRACT).json()
    assert next(v for v in body["variants"] if v["variant_id"] == "oat-linen")["price"]["amount"] == 124900
    # "hidden" takes a variant out of a region.
    put("AE", [{"sku": "SOFA-OSLO-3", "variant_id": "charcoal-wool", "price": "1.00", "status": "hidden"}])
    ae = client.get("/supplier/prices", params={"region": "AE"}, headers=owner).json()["rows"]
    assert next(r for r in ae if r["variant_id"] == "charcoal-wool")["sold"] is False
    # Regions: delivery defaults, validated with the currency's decimals.
    error_of(client.put("/supplier/regions", json={"regions": [{"region": "GB", "default_delivery_fee": "49.001"}]}, headers=owner), 422, "validation_failed")
    res = client.put("/supplier/regions", json={"regions": [{"region": "GB", "default_delivery_fee": "49.00", "delivery_days_min": 5, "delivery_days_max": 9}]}, headers=owner)
    gb = next(r for r in res.json()["regions"] if r["region"] == "GB")
    assert (gb["default_delivery_fee"], gb["served"]) == ("49.00", True)
    # A viewer cannot save prices.
    viewer = as_supplier(_join(client, owner, "viewer@example.com", "viewer"), SUPPLIER_ID)
    error_of(put_as(client, viewer), 403, "forbidden")


def put_as(client, headers):
    return client.put("/supplier/prices", params={"region": "GB"}, json={"rows": [{"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1.00"}]}, headers=headers)


# --- Spreadsheet imports ----------------------------------------------------------------------------


def _xlsx(rows: list[dict], *, numbers: bool = True, sheet: str = "catalogue", workbook: bytes | None = None) -> bytes:
    """The rows in a workbook; with `numbers`, cells a spreadsheet would hold
    as numbers are numbers (dimensions, stock, days, the chair's price)."""
    if workbook:
        wb = load_workbook(io.BytesIO(workbook))
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        header = list(importer.COLUMNS)
        ws.append(header)
    ints = {"width_mm", "height_mm", "depth_mm", "delivery_days_min", "delivery_days_max", "stock", "tax_rate_percent"}
    for r in rows:
        values = []
        for col in header:
            v = r.get(col, "")
            if numbers and col in ints and v not in ("", None):
                v = int(v)
            elif numbers and col == "price" and r["sku"] == "CHAIR-BERGEN" and v:
                v = float(v)  # 649.0: read back through Decimal(str(cell))
            values.append(v if v != "" else None)
        ws.append(values)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload(client, headers, data: bytes, name: str = "catalogue.xlsx", mode: str = "upsert"):
    return client.post("/supplier/imports", data={"mode": mode}, files={"file": (name, data)}, headers=headers)


def test_supplier_xlsx_dry_run_and_apply(client, monkeypatch):
    no_internet(monkeypatch)
    owner, sid = verified(client)
    rows = rows_of(FEED_CSV)
    res = _upload(client, owner, _xlsx(rows))
    assert res.status_code == 201, res.text
    imp = res.json()
    dry = imp["dry_run"]
    assert (dry["rows"], dry["created"], dry["updated"], dry["unchanged"], dry["rejected"]) == (6, 6, 0, 0, 0)
    assert (dry["new_products"], dry["new_variants"]) == (2, 3) and imp["format"] == "xlsx" and imp["expires_at"]
    with SessionLocal() as db:
        assert db.scalars(select(Product).where(Product.supplier_id == sid)).first() is None  # a dry run writes nothing
    # 1299.999 in a spreadsheet cell (a number) is rejected as bad_price, with its row.
    bad = [dict(r) for r in rows]
    bad[0]["price"] = "1299.999"
    data = _xlsx(bad)
    wb = load_workbook(io.BytesIO(data))
    wb["catalogue"]["Q2"] = 1299.999
    out = io.BytesIO()
    wb.save(out)
    res = _upload(client, owner, out.getvalue())
    dry_bad = res.json()["dry_run"]
    assert dry_bad["rejected"] == 1 and dry_bad["errors"][0] == {"row": 2, "column": "price", "code": "bad_price", "value": "1299.999"}
    # Apply the good one: a feed_runs row with the same counts.
    res = client.post(f"/supplier/imports/{imp['import_id']}/apply", headers=owner)
    assert res.status_code == 202, res.text
    run = client.get(f"/supplier/imports/{imp['import_id']}", headers=owner).json()["run"]
    assert run["state"] == "done" and run["source"] == "portal"
    assert (run["rows"], run["created"], run["updated"], run["unchanged"], run["rejected"]) == (6, 6, 0, 0, 0)
    with SessionLocal() as db:
        fr = db.get(FeedRun, run["feed_id"])
        assert fr.supplier_id == sid and fr.format == "csv"
        assert {p.status for p in db.scalars(select(Product).where(Product.supplier_id == sid))} == {"pending_review"}
    error_of(client.post(f"/supplier/imports/{imp['import_id']}/apply", headers=owner), 409, "already_applied")
    history = client.get("/supplier/imports", headers=owner).json()
    assert history["runs"][0]["feed_id"] == run["feed_id"] and len(history["imports"]) == 2
    assert client.get(f"/supplier/runs/{run['feed_id']}", headers=owner).json()["created"] == 6
    # CSV and JSON uploads use the same dry run; a wrong header is 422 feed_invalid.
    assert _upload(client, owner, FEED_CSV, "feed.csv").json()["dry_run"]["unchanged"] == 6
    error_of(_upload(client, owner, b"a,b\r\n1,2\r\n", "feed.csv"), 422, "feed_invalid")
    error_of(_upload(client, owner, b"not a workbook", "feed.xlsx"), 422, "feed_invalid")
    error_of(_upload(client, owner, FEED_CSV, "feed.txt"), 422, "validation_failed")
    # Dry runs are kept 24 h.
    clk = Clock(monkeypatch)
    stale = _upload(client, owner, FEED_CSV, "feed.csv").json()["import_id"]
    clk.advance(hours=25)
    error_of(client.post(f"/supplier/imports/{stale}/apply", headers=owner), 410, "expired")
    with SessionLocal() as db:
        assert imports.purge(db, clk.at) >= 1
        assert db.get(SupplierImport, stale) is None
        assert db.get(SupplierImport, imp["import_id"]) is not None  # applied ones stay


def test_supplier_template_round_trip(client, monkeypatch):
    no_internet(monkeypatch)
    owner, sid = verified(client)
    res = client.get("/supplier/imports/template.xlsx", headers=owner)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["catalogue", "help", "categories", "regions"]
    ws = wb["catalogue"]
    assert [c.value for c in ws[1]] == list(importer.COLUMNS)
    assert ws.column_dimensions["A"].number_format == "@"  # sku as text
    lists = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert any(f.startswith("categories!") for f in lists) and any(f.startswith("regions!") for f in lists)
    cats = [r[0] for r in wb["categories"].iter_rows(min_row=2, values_only=True)]
    assert "furniture/seating/sofas" in cats
    regions = list(wb["regions"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in regions[:2]] == ["GB", "AE"] and regions[0][7] == "yes"  # the supplier's regions first
    assert any(r[0] == "price" for r in wb["help"].iter_rows(min_row=2, values_only=True))
    # Filled with the fixture's rows (as a spreadsheet would hold them), it imports without errors.
    filled = _xlsx(rows_of(FEED_CSV), workbook=res.content)
    up = _upload(client, owner, filled)
    assert up.status_code == 201, up.text
    assert up.json()["dry_run"]["rejected"] == 0, up.json()["dry_run"]["errors"]
    client.post(f"/supplier/imports/{up.json()['import_id']}/apply", headers=owner)
    run = client.get(f"/supplier/imports/{up.json()['import_id']}", headers=owner).json()["run"]
    assert (run["state"], run["rejected"], run["created"]) == ("done", 0, 6)
    # A SKU a spreadsheet turned into a number is imported but reported.
    wb = load_workbook(io.BytesIO(filled))
    wb["catalogue"]["A2"] = 123
    out = io.BytesIO()
    wb.save(out)
    warnings = _upload(client, owner, out.getvalue()).json()["dry_run"]["warnings"]
    assert {"row": 2, "column": "sku", "code": "looks_reformatted", "value": "123"} in warnings


# --- Inbox -------------------------------------------------------------------------------------------


def test_supplier_inbox_quote_and_order_states(client, monkeypatch):
    owner = _fixture_supplier(client)
    buyer = signup(client, email="buyer@example.com")
    quote = client.post("/market/orders", json=order_body("quote"), headers={**buyer, **CONTRACT}).json()
    assert quote["state"] == "submitted"
    mails = outbox("New request for quote")
    assert mails and "Layla Haddad" in mails[0].text and "After 1 Dec" in mails[0].text
    # The inbox shows the lead with the buyer's contact (the supplier's own lead).
    items = client.get("/supplier/inbox", headers=owner).json()
    assert items["open"] == {"requests": 1, "orders": 0}
    item = items["items"][0]
    assert item["kind"] == "quote" and item["buyer"]["email"] == "buyer@example.com" and item["open"]
    # A catalogue member gets 403; an orders member answers.
    catalogue = as_supplier(_join(client, owner, "cat@example.com", "catalogue"), SUPPLIER_ID)
    error_of(client.get("/supplier/inbox", headers=catalogue), 403, "forbidden")
    error_of(client.post(f"/supplier/inbox/{quote['order_id']}/quote", json={}, headers=catalogue), 403, "forbidden")
    orders = as_supplier(_join(client, owner, "orders@example.com", "orders"), SUPPLIER_ID)
    error_of(client.post(f"/supplier/inbox/{quote['order_id']}/quote", json={"lines": [{"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1180.001"}]}, headers=orders), 422, "validation_failed")
    res = client.post(
        f"/supplier/inbox/{quote['order_id']}/quote",
        json={"lines": [{"sku": "SOFA-OSLO-3", "variant_id": "oat-linen", "price": "1180.00"}], "delivery_fee": "150.00", "message": "Assembled in Dubai."},
        headers=orders,
    )
    assert res.status_code == 200, res.text
    assert res.json()["state"] == "quoted" and res.json()["lines"][0]["quoted_unit"]["amount"] == 118000
    seen = client.get(f"/market/orders/{quote['order_id']}", headers={**buyer, **CONTRACT}).json()
    assert seen["state"] == "quoted" and seen["lines"][0]["quoted_price"]["amount"] == 118000
    assert seen["suppliers"][0]["quote"]["message"] == "Assembled in Dubai."
    # The buyer accepts: payments are off, so the order is the supplier's to fulfil now.
    order = client.post(f"/market/orders/{quote['order_id']}/accept", headers=buyer).json()
    assert order["state"] == "accepted" and order["payment"] == "offline"
    assert outbox("New order")
    error_of(client.post(f"/supplier/inbox/{order['order_id']}/ship", json={"carrier": ""}, headers=orders), 422, "validation_failed")
    res = client.post(f"/supplier/inbox/{order['order_id']}/ship", json={"carrier": "DHL", "reference": "JD0123"}, headers=orders)
    assert res.json()["state"] == "shipped" and res.json()["shipment"] == {"carrier": "DHL", "reference": "JD0123"}
    seen = client.get(f"/market/orders/{order['order_id']}", headers={**buyer, **CONTRACT}).json()
    assert seen["state"] == "shipped" and seen["suppliers"][0]["shipment"]["reference"] == "JD0123"
    assert client.post(f"/supplier/inbox/{order['order_id']}/deliver", headers=orders).json()["state"] == "delivered"
    assert client.get(f"/market/orders/{order['order_id']}", headers={**buyer, **CONTRACT}).json()["state"] == "delivered"
    error_of(client.post(f"/supplier/inbox/{order['order_id']}/accept", headers=orders), 409, "conflict")

    # A paid order (payments on): accept or reject, after the payment only.
    with SessionLocal() as db:
        seed.load_fixture(db, FIXTURES, payments_ready=True)
    payments_on(monkeypatch)
    fake = FakeGateway()
    from app.market import checkout

    monkeypatch.setattr(checkout, "gateway", lambda: fake)
    placed = client.post("/market/orders", json=order_body(), headers={**buyer, **CONTRACT}).json()
    error_of(client.get(f"/supplier/inbox/{placed['order_id']}", headers=orders), 404, "not_found")  # not paid yet
    url = client.post(f"/market/checkout/{placed['order_id']}", headers=buyer).json()["url"]
    session_id = url.rsplit("/", 1)[1]
    event = {"id": "evt_1", "type": "checkout.session.completed", "data": {"object": {
        "id": session_id, "payment_status": "paid", "payment_intent": "pi_1", "client_reference_id": placed["order_id"],
        "metadata": {"kind": "market_order", "order_id": placed["order_id"]}}}}  # fmt: skip
    assert stripe_post(client, event).status_code == 200
    assert len({m.subject for m in outbox("New order")}) == 2
    assert client.get(f"/supplier/inbox/{placed['order_id']}", headers=orders).json()["state"] == "pending"
    assert client.post(f"/supplier/inbox/{placed['order_id']}/accept", headers=orders).json()["state"] == "accepted"
    assert client.get(f"/market/orders/{placed['order_id']}", headers={**buyer, **CONTRACT}).json()["state"] == "accepted"
    # The buyer cancels a request: the supplier is told.
    second = client.post("/market/orders", json=order_body("quote"), headers={**buyer, **CONTRACT}).json()
    client.post(f"/market/orders/{second['order_id']}/cancel", headers={**buyer, **CONTRACT})
    assert outbox("Cancelled by the buyer: request for quote")
    # Another supplier never sees these leads.
    stranger, _ = verified(client, email="other@example.com", name="Other Co")
    assert client.get("/supplier/inbox", headers=stranger).json()["items"] == []
    error_of(client.get(f"/supplier/inbox/{quote['order_id']}", headers=stranger), 404, "not_found")


# --- Analytics ------------------------------------------------------------------------------------------


def test_supplier_analytics_counts_only(client):
    owner = _fixture_supplier(client)
    buyer = signup(client, email="buyer@example.com")
    sofa = SOFA["product_id"]
    assert client.get("/market/search", params={"q": "sofa", "region": "AE"}, headers=CONTRACT).status_code == 200
    client.get("/market/search", params={"q": "sofa", "region": "AE"}, headers=CONTRACT)
    product = client.get(f"/market/products/{sofa}", params={"region": "AE"}, headers=CONTRACT).json()
    geometry = next(v["geometry"] for v in product["variants"] if v["geometry"])
    assert "/market/geometry/" in geometry["url"] and "region=AE" in geometry["url"]
    res = client.get(geometry["url"].replace("https://api.truebex.com", ""), follow_redirects=False)
    assert res.status_code == 302 and "/files/market/geometry/" in res.headers["location"]
    client.post("/market/orders", json=order_body("quote"), headers={**buyer, **CONTRACT})

    body = client.get("/supplier/analytics", params={"region": "AE"}, headers=owner).json()
    row = next(p for p in body["products"] if p["product_id"] == sofa)
    assert (row["impressions"], row["views"], row["geometry_downloads"], row["quotes"], row["orders"]) == (2, 1, 1, 1, 0)
    assert len(body["days"]) == 30 and body["days"][-1]["views"] == 1 and body["totals"]["quotes"] == 1
    assert body["region"] == "AE"
    # Counts only: no buyer, account or contact anywhere in the answer.
    text = str(body)
    assert "buyer@example.com" not in text and "Layla" not in text and "user_id" not in text
    with SessionLocal() as db:
        columns = set(MarketEventDaily.__table__.columns.keys())
    assert columns == {"id", "product_id", "supplier_id", "region", "day", "impressions", "views", "geometry_downloads", "quotes", "orders", "order_value"}
    # An accepted quote counts as an order with its value (AED, minor units).
    with SessionLocal() as db:
        from app.market import orders as m_orders
        from app.market.models import MarketOrder
        from app.market.schemas import SupplierQuoteIn

        q = db.scalar(select(MarketOrder).where(MarketOrder.kind == "quote"))
        m_orders.supplier_quote(db, q, SUPPLIER_ID, SupplierQuoteIn(lines=[]))
    quote_id = client.get("/market/orders", headers={**buyer, **CONTRACT}).json()["orders"][0]["order_id"]
    client.post(f"/market/orders/{quote_id}/accept", headers=buyer)
    body = client.get("/supplier/analytics", params={"region": "AE"}, headers=owner).json()
    assert body["totals"]["orders"] == 1 and body["totals"]["order_value"] == [{"amount": 129900, "currency": "AED", "exponent": 2}]
    # Validation and auth.
    error_of(client.get("/supplier/analytics", params={"from": "2026-13-01"}, headers=owner), 422, "validation_failed")
    error_of(client.get("/supplier/analytics", params={"from": "2026-10-09", "to": "2026-10-01"}, headers=owner), 422, "validation_failed")
    error_of(client.get("/supplier/analytics", params={"region": "XX"}, headers=owner), 422, "validation_failed")
    error_of(client.get("/supplier/analytics"), 401, "unauthenticated")


# --- Listing plan and billing (PF2's interface, mocked) --------------------------------------------------


class _FakeProvider:
    def __init__(self):
        self.checkouts = []
        self.invoices = []

    def enabled(self) -> bool:
        return True

    def create_checkout(self, db, user, payment, price, seats, discount):
        self.checkouts.append((user.email, payment.plan, payment.amount, payment.currency, price.provider_price_id, seats, discount))
        return f"https://checkout.stripe.test/{payment.reference}"

    def create_invoice(self, db, customer, lines):
        self.invoices.append((customer, [(ln.description, ln.amount_minor, ln.currency) for ln in lines]))
        return "in_test_1"

    def invoice_pdf_url(self, db, user, invoice_id):
        return f"https://invoices.stripe.test/{invoice_id}.pdf"


def _pf2(monkeypatch) -> _FakeProvider:
    """Stand in for PF2's billing interface: billing.base and get_provider."""
    from app.billing import providers

    provider = _FakeProvider()

    class InvoiceLine:
        def __init__(self, description, amount_minor, currency, quantity=1):
            self.description, self.amount_minor, self.currency, self.quantity = description, amount_minor, currency, quantity

    base = types.ModuleType("app.billing.base")
    base.InvoiceLine = InvoiceLine
    base.ProviderError = type("ProviderError", (Exception,), {})
    monkeypatch.setitem(sys.modules, "app.billing.base", base)
    monkeypatch.setattr(providers, "get_provider", lambda name, settings=None: provider, raising=False)
    monkeypatch.setattr(
        listing, "listing_price", lambda db, plan, currency, amount: types.SimpleNamespace(provider_price_id=f"price_{plan}_{currency}_{amount}")
    )
    return provider


def test_supplier_listing_checkout_via_billing(client, monkeypatch):
    owner, sid = verified(client)
    from app.market import listing as m_listing

    plans = dict(m_listing.plans())
    paid = m_listing.ListingPlan("standard", "Standard", 800, {"GBP": 4900})
    monkeypatch.setattr(m_listing, "plans", lambda: {**plans, "standard": paid})
    monkeypatch.setattr(m_listing, "plan_of", lambda s: {**plans, "standard": paid}.get(s.listing_plan or "founding"))

    view = client.get("/supplier/listing", headers=owner).json()
    assert view["plan"] == "founding" and view["currency"] == "GBP"
    assert next(p for p in view["plans"] if p["id"] == "standard")["monthly_fee"] == {"amount": 4900, "currency": "GBP", "exponent": 2}
    assert view["billing"] == {"interface": True, "stripe_ready": True}
    # PF2 is in, but no listing price is synced yet (GD7 sets the fees): the paid plan waits as pending.
    res = client.post("/supplier/listing", json={"plan": "standard"}, headers=owner).json()
    assert res["state"] == "pending" and res["checkout_url"] is None
    assert "isn't set up yet" in res["detail"]
    error_of(client.post("/supplier/listing", json={"plan": "gold"}, headers=owner), 422, "validation_failed")

    # Through PF2's interface (mocked): a checkout on the synced listing price.
    provider = _pf2(monkeypatch)
    res = client.post("/supplier/listing", json={"plan": "standard"}, headers=owner)
    assert res.status_code == 200, res.text
    assert res.json()["state"] == "checkout" and res.json()["checkout_url"].startswith("https://checkout.stripe.test/lst_")
    assert provider.checkouts == [("owner@nord.example.com", "listing_standard", 4900, "GBP", "price_standard_GBP_4900", 1, None)]
    with SessionLocal() as db:
        row = db.scalar(select(ListingSubscription).where(ListingSubscription.status == "checkout"))
        assert (row.supplier_id, row.plan, row.provider, row.amount) == (sid, "standard", "stripe", 4900)
        payment = db.scalar(select(Payment).where(Payment.reference == row.payment_ref))
        assert payment.plan == "listing_standard" and payment.status == "pending"
        # PF2's webhook writes the subscription; the sync job follows it.
        db.add(Subscription(user_id=payment.user_id, plan="listing_standard", provider="stripe", status="active",
                            provider_customer_id="cus_nord", provider_subscription_id="sub_nord"))  # fmt: skip
        db.commit()
        assert listing.sync(db, row.created_at) == 1
        from app.market.models import Supplier

        supplier = db.get(Supplier, sid)
        assert (supplier.listing_plan, supplier.billing_customer_id) == ("standard", "cus_nord")
        assert db.scalar(select(ListingSubscription).where(ListingSubscription.payment_ref == row.payment_ref)).status == "active"
        assert listing.billed_by_subscription(db, supplier)
        # The live listing subscription is not the owner's app plan (PF2 skips listing_* tiers).
        from app.billing import service as billing_service

        buyer = db.get(User, payment.user_id)
        assert billing_service.live_subscription(db, buyer) is None
        assert billing_service.managed_subscription(db, buyer) is None
        assert billing_service.effective_plan(db, buyer) == "free"
    view = client.get("/supplier/listing", headers=owner).json()
    assert view["plan"] == "standard" and view["subscription"]["status"] == "active" and view["commission_bp"] == 800

    # Statements: the commission invoiced through create_invoice; the fee is the subscription's.
    with SessionLocal() as db:
        from app.market.common import now

        db.add(Commission(order_id="a" * 32, supplier_id=sid, rate_bp=800, base=108250, amount=8660, currency="GBP", exponent=2,
                          state="accrued", created_at=now() - timedelta(days=40)))  # fmt: skip
        db.commit()
        made = commissions.run_statements(db, now())
        assert len(made) == 1 and made[0].state == "invoiced" and made[0].listing_fee == 0
    assert provider.invoices == [("cus_nord", [(provider.invoices[0][1][0][0], 8660, "gbp")])]
    statements = client.get("/supplier/statements", headers=owner).json()
    assert statements["statements"][0]["invoice_ref"] == "in_test_1" and statements["commissions"][0]["amount"]["amount"] == 8660
    statement_id = statements["statements"][0]["statement_id"]
    assert client.get(f"/supplier/statements/{statement_id}/invoice", headers=owner).json()["url"].endswith("in_test_1.pdf")
    # Payouts: Stripe Connect onboarding (PF7's gateway, faked).
    from app.market import checkout

    monkeypatch.setattr(checkout, "gateway", lambda: FakeGateway())
    assert client.post("/supplier/payouts/connect", headers=owner).json()["url"].startswith("https://connect.stripe.test/")
    # A plan without a fee applies at once.
    assert client.post("/supplier/listing", json={"plan": "founding"}, headers=owner).json()["state"] == "active"
    # Owners only.
    cat = as_supplier(_join(client, owner, "cat@example.com", "catalogue"), sid)
    error_of(client.get("/supplier/listing", headers=cat), 403, "forbidden")
    error_of(client.get("/supplier/statements", headers=cat), 403, "forbidden")


# --- Team and keys ----------------------------------------------------------------------------------------


def test_supplier_members_and_scoped_keys(client, monkeypatch):
    no_internet(monkeypatch)
    owner, sid = verified(client)
    error_of(client.post("/supplier/members", json={"email": "bad", "role": "viewer"}, headers=owner), 422, "validation_failed")
    error_of(client.post("/supplier/members", json={"email": "a@example.com", "role": "admin"}, headers=owner), 422, "validation_failed")
    res = client.post("/supplier/members", json={"email": "Cat@Example.com", "role": "catalogue"}, headers=owner)
    assert res.status_code == 201 and res.json()["role"] == "catalogue"
    invite = outbox("invited you to Nord Living")[-1]
    assert invite.to.lower() == "cat@example.com" and "catalogue editor" in invite.text
    token = invite.text.split("token=")[1].split()[0]
    # The wrong account cannot accept; the invited address can (any case).
    wrong = signup(client, email="someone@example.com")
    error_of(client.post("/supplier/invites/accept", json={"token": token}, headers=wrong), 403, "wrong_account")
    error_of(client.post("/supplier/invites/accept", json={"token": "x" * 43}, headers=wrong), 404, "not_found")
    cat = signup(client, email="cat@example.com")
    assert client.post("/supplier/invites/accept", json={"token": token}, headers=cat).json()["role"] == "catalogue"
    team = client.get("/supplier/members", headers=owner).json()
    assert {(m["email"], m["role"]) for m in team["members"]} == {("owner@nord.example.com", "owner"), ("cat@example.com", "catalogue")}
    assert team["invites"] == []
    error_of(client.get("/supplier/members", headers=as_supplier(cat, sid)), 403, "forbidden")
    error_of(client.post("/supplier/members", json={"email": "cat@example.com", "role": "viewer"}, headers=owner), 409, "already_member")
    # An expired invitation.
    clk = Clock(monkeypatch)
    client.post("/supplier/members", json={"email": "late@example.com", "role": "viewer"}, headers=owner)
    late_token = outbox("invited you", to="late@example.com")[-1].text.split("token=")[1].split()[0]
    clk.advance(days=8)
    late = signup(client, email="late@example.com")
    error_of(client.post("/supplier/invites/accept", json={"token": late_token}, headers=late), 410, "invite_expired")

    # A key made in the portal carries supplier_id, is shown once, and works on 5.11.
    res = client.post("/supplier/keys", json={"name": "ERP"}, headers=owner)
    assert res.status_code == 201 and res.json()["key"].startswith("tbx_live_")
    key = res.json()["key"]
    with SessionLocal() as db:
        from app.models import ApiKey

        assert db.scalar(select(ApiKey.supplier_id).where(ApiKey.prefix == res.json()["prefix"])) == sid
    listed = client.get("/supplier/members", headers=owner).json()["keys"]
    assert listed[0]["name"] == "ERP" and "key" not in listed[0]
    files = {"file": ("feed.csv", FEED_CSV, "text/csv")}
    up = client.post("/market/feeds", data={"format": "csv"}, files=files, headers={"X-API-Key": key, **CONTRACT})
    assert up.status_code == 202, up.text
    # Developer keys and supplier keys stay apart.
    assert client.get("/keys", headers=owner).json() == []
    # A member who leaves loses the keys they made; the last owner stays.
    cat_key = None
    client.patch(f"/supplier/members/{_uid('cat@example.com')}", json={"role": "owner"}, headers=owner)
    cat_key = supplier_key(client, as_supplier(cat, sid), "cat's key")
    client.delete(f"/supplier/members/{_uid('cat@example.com')}", headers=owner)
    error_of(client.post("/market/feeds", data={"format": "csv"}, files=files, headers={"X-API-Key": cat_key, **CONTRACT}), 401, "unauthenticated")
    error_of(client.delete(f"/supplier/members/{_uid('owner@nord.example.com')}", headers=owner), 409, "last_owner")
    error_of(client.patch(f"/supplier/members/{_uid('owner@nord.example.com')}", json={"role": "viewer"}, headers=owner), 409, "last_owner")
    # Revoke: the key stops working.
    key_id = listed[0]["id"]
    assert client.delete(f"/supplier/keys/{key_id}", headers=owner).json()["revoked_at"]
    error_of(client.post("/market/feeds", data={"format": "csv"}, files=files, headers={"X-API-Key": key, **CONTRACT}), 401, "unauthenticated")


def _uid(email: str) -> int:
    with SessionLocal() as db:
        return db.scalar(select(User.id).where(User.email == email))


# --- Every route: auth ------------------------------------------------------------------------------------

ROUTES = [
    ("GET", "/supplier/overview"),
    ("GET", "/supplier/products"),
    ("POST", "/supplier/products"),
    ("POST", "/supplier/products/submit"),
    ("GET", f"/supplier/products/{'a' * 32}"),
    ("PATCH", f"/supplier/products/{'a' * 32}"),
    ("POST", f"/supplier/products/{'a' * 32}/variants"),
    ("PATCH", f"/supplier/products/{'a' * 32}/variants/v"),
    ("DELETE", f"/supplier/products/{'a' * 32}/variants/v"),
    ("POST", f"/supplier/products/{'a' * 32}/variants/v/images"),
    ("DELETE", f"/supplier/products/{'a' * 32}/variants/v/images/{'b' * 64}"),
    ("POST", f"/supplier/products/{'a' * 32}/submit"),
    ("POST", "/supplier/geometry"),
    ("POST", "/supplier/documents"),
    ("GET", "/supplier/prices"),
    ("PUT", "/supplier/prices"),
    ("GET", "/supplier/regions"),
    ("PUT", "/supplier/regions"),
    ("GET", "/supplier/imports"),
    ("POST", "/supplier/imports"),
    ("GET", "/supplier/imports/template.xlsx"),
    ("GET", f"/supplier/imports/{'a' * 32}"),
    ("POST", f"/supplier/imports/{'a' * 32}/apply"),
    ("GET", f"/supplier/runs/{'a' * 32}"),
    ("GET", "/supplier/feeds/source"),
    ("PUT", "/supplier/feeds/source"),
    ("DELETE", "/supplier/feeds/source"),
    ("POST", "/supplier/feeds/source/pull"),
    ("GET", "/supplier/inbox"),
    ("GET", f"/supplier/inbox/{'a' * 32}"),
    ("POST", f"/supplier/inbox/{'a' * 32}/quote"),
    ("POST", f"/supplier/inbox/{'a' * 32}/accept"),
    ("POST", f"/supplier/inbox/{'a' * 32}/ship"),
    ("GET", "/supplier/analytics"),
    ("GET", "/supplier/listing"),
    ("POST", "/supplier/listing"),
    ("GET", "/supplier/statements"),
    ("GET", f"/supplier/statements/{'a' * 32}/invoice"),
    ("POST", "/supplier/payouts/connect"),
    ("GET", "/supplier/members"),
    ("POST", "/supplier/members"),
    ("PATCH", "/supplier/members/1"),
    ("DELETE", "/supplier/members/1"),
    ("DELETE", f"/supplier/invites/{'a' * 32}"),
    ("POST", "/supplier/keys"),
    ("DELETE", "/supplier/keys/1"),
]
SESSION_ONLY = [("GET", "/supplier/me"), ("POST", "/supplier/applications"), ("POST", "/supplier/invites/accept")]


def test_supplier_routes_need_auth(client):
    """Every portal route: 401 without a session (the envelope), 403
    `not_supplier` for an account that belongs to no supplier."""
    stranger = signup(client, email="stranger@example.com")
    for method, path in ROUTES + SESSION_ONLY:
        res = client.request(method, path)
        error_of(res, 401, "unauthenticated")
    for method, path in ROUTES:
        res = client.request(method, path, headers=stranger, json={})
        assert res.status_code in (403, 422), (method, path, res.text)
        if res.status_code == 403:
            assert res.json()["code"] == "not_supplier", (method, path)
    # Naming a supplier one does not belong to is the same as none.
    owner, sid = verified(client)
    error_of(client.get("/supplier/overview", headers=as_supplier(stranger, sid)), 403, "not_supplier")
    assert client.get("/supplier/overview", headers=owner).json()["supplier"]["supplier_id"] == sid


def test_supplier_overview_and_switcher(client, monkeypatch):
    no_internet(monkeypatch)
    owner, sid = verified(client)
    over = client.get("/supplier/overview", headers=owner).json()
    assert over["role"] == "owner" and over["open"] == {"requests": 0, "orders": 0} and over["last_run"] is None
    assert over["week"] == {"impressions": 0, "views": 0, "geometry_downloads": 0, "quotes": 0, "orders": 0}
    # A member of two suppliers switches with the header.
    from app.market.catalogue import add_member, create_supplier

    with SessionLocal() as db:
        other = create_supplier(db, name="Alpha Co", country="AE", status="verified")
        add_member(db, other, db.scalar(select(User).where(User.email == "owner@nord.example.com")), "viewer")
        db.commit()
        other_id = other.supplier_id
    plain = {"Authorization": owner["Authorization"]}
    me = client.get("/supplier/me", headers=plain).json()
    assert [m["name"] for m in me["memberships"]] == ["Alpha Co", "Nord Living"]
    assert me["supplier"]["name"] == "Alpha Co" and me["role"] == "viewer"  # the first by name without a header
    me = client.get("/supplier/me", headers=as_supplier(plain, sid)).json()
    assert me["supplier"]["supplier_id"] == sid and me["role"] == "owner"
    error_of(client.post("/supplier/products", json={"sku": "S", "name": "S", "category": "furniture/seating"}, headers=as_supplier(plain, other_id)), 403, "forbidden")
