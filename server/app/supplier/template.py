"""The spreadsheet template (`GET /supplier/imports/template.xlsx`): the
contract's §6.4 columns in one sheet named `catalogue`, one row per variant
and region, plus `help` (every column with its rule and an example, GD1
§6.3), `categories` and `regions` (the lists the drop-downs use, the
supplier's regions first).

`sku`, `variant_id`, `gtin`, `price`, `delivery_fee` and
`tax_rate_percent` are formatted as text, so a spreadsheet program keeps
`00123`, a 13-digit GTIN and `1299.00` as typed.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from ..market.importer import COLUMNS, REQUIRED_HEADER
from ..market.models import Supplier
from ..market.taxonomy import active_categories, active_regions
from .prices import _major, _pct, served

TEXT_COLUMNS = ("sku", "variant_id", "gtin", "price", "delivery_fee", "tax_rate_percent")
DATA_ROWS = 20_000  # rows the drop-downs and text formats cover

# (column, level, required, rule, example) — GD1 §6.3, the contract §6.4.
HELP = (
    ("sku", "product", "yes", "≤ 64 characters; stable across uploads, never reused for another product", "SOFA-OSLO-3"),
    ("variant_id", "variant", "yes", "≤ 64; stable; default for a product with one variant", "oat-linen"),
    ("name", "product", "yes", "≤ 120, without the variant (the options say that)", "Oslo 3-seater sofa"),
    ("category", "product", "yes", "a path from the categories sheet", "furniture/seating/sofas"),
    ("kind", "product", "no (object)", "object, material, finish or theme", "object"),
    ("option_size", "variant", "no", "≤ 40", "3-seater"),
    ("option_colour", "variant", "no", "≤ 40", "Oat"),
    ("option_finish", "variant", "no", "≤ 40", "Linen"),
    ("materials", "variant", "no", "lower-case words separated by ; (≤ 40 each)", "linen;oak"),
    ("width_mm", "variant", "for objects", "whole millimetres, side to side as you face the front", "2100"),
    ("height_mm", "variant", "for objects", "whole millimetres, floor to top", "850"),
    ("depth_mm", "variant", "for objects", "whole millimetres, front to back", "950"),
    ("geometry_url", "variant", "for objects", "https link to one 3D file per variant: GLB (preferred), glTF, FBX, OBJ or .tbxa", "https://cdn.example.com/oslo/oat-linen.glb"),
    ("image_urls", "variant", "yes", "https links separated by ;, JPEG or PNG, at least 512 px on the short side; the first is the thumbnail", "https://cdn.example.com/oslo/oat-linen-1.jpg"),
    ("region", "row", "yes", "a region from the regions sheet", "GB"),
    ("currency", "row", "yes", "the region's currency (regions sheet)", "GBP"),
    ("price", "row", "yes", "decimal text in the main unit with at most the currency's decimals; no thousands separator, no symbol", "1299.00"),
    ("price_includes_tax", "row", "yes", "true or false", "true"),
    ("tax_rate_percent", "row", "no (the region's rate)", "percent", "20"),
    ("delivery_fee", "row", "no (your region default)", "decimal text, as price", "49.00"),
    ("delivery_days_min", "row", "no", "whole days from acceptance to delivery", "7"),
    ("delivery_days_max", "row", "no", "whole days", "14"),
    ("stock", "row", "no", "whole number; empty = not tracked", "12"),
    ("lead_time_days", "row", "no", "whole days, for made-to-order products", ""),
    ("status", "row", "no (active)", "active, discontinued or hidden (not sold in this region)", "active"),
    ("description", "product", "no", "plain text, ≤ 2000", "Three-seat sofa on a solid oak frame, with removable covers."),
    ("brand", "product", "no", "≤ 70", "Nord Living"),
    ("gtin", "variant", "no", "GS1 GTIN-8, -12, -13 or -14 as digits, valid check digit", "9501101530003"),
    ("classification", "product", "no", "scheme:code pairs separated by ; (uniclass, etim, gpc)", "uniclass:Pr_40_50_12_81"),
    ("availability", "row", "no", "in_stock, low_stock, made_to_order (needs lead_time_days) or out_of_stock; overrides the stock rule", ""),
)
NOTES = (
    "One row per variant per region: a product with 2 variants sold in GB and AE takes 4 rows.",
    "Product columns repeat on every row of a SKU and must agree; variant columns on every row of a variant.",
    "Availability: discontinued wins; then the availability column; then stock above 0 = in stock, 0 = out of stock "
    "(made to order with a lead time), empty = in stock and not tracked (made to order with a lead time).",
    "Upload the file in the portal under Imports: a dry run shows what would be created, updated, left unchanged and "
    "rejected (row, column and code) before you apply it.",
)

_BOLD = Font(bold=True)
_HEAD = PatternFill("solid", fgColor="DDEBFF")


def _header(ws, names) -> None:
    ws.append(list(names))
    for cell in ws[1]:
        cell.font = _BOLD
        cell.fill = _HEAD
    ws.freeze_panes = "A2"


def _list(formula: str) -> DataValidation:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
    dv.errorTitle, dv.error = "Not in the list", "Pick a value from the list."
    return dv


def build(db: Session, supplier: Supplier) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "catalogue"
    _header(ws, COLUMNS)
    letter = {name: ws.cell(row=1, column=i + 1).column_letter for i, name in enumerate(COLUMNS)}
    for name in COLUMNS:
        dim = ws.column_dimensions[letter[name]]
        dim.width = 28 if name in ("name", "category", "geometry_url", "image_urls", "description") else 16
        if name in TEXT_COLUMNS:
            dim.number_format = "@"
        if name in REQUIRED_HEADER:
            ws.cell(row=1, column=COLUMNS.index(name) + 1).font = Font(bold=True, color="1762C9")

    cats = list(active_categories(db).values())
    regions = active_regions(db)
    mine = {k: v for k, v in served(db, supplier).items() if v.active}
    ordered = [r for k, r in regions.items() if k in mine] + [r for k, r in regions.items() if k not in mine]

    def validate(column: str, dv: DataValidation) -> None:
        dv.add(f"{letter[column]}2:{letter[column]}{DATA_ROWS + 1}")
        ws.add_data_validation(dv)

    validate("category", _list(f"categories!$A$2:$A${len(cats) + 1}"))
    validate("region", _list(f"regions!$A$2:$A${len(ordered) + 1}"))
    validate("currency", _list(f"regions!$C$2:$C${len(ordered) + 1}"))
    validate("kind", _list('"object,material,finish,theme"'))
    validate("price_includes_tax", _list('"true,false"'))
    validate("status", _list('"active,discontinued,hidden"'))
    validate("availability", _list('"in_stock,low_stock,made_to_order,out_of_stock"'))

    help_ws = wb.create_sheet("help")
    _header(help_ws, ("column", "level", "required", "rule", "example"))
    for row in HELP:
        help_ws.append(list(row))
    help_ws.append([])
    for note in NOTES:
        help_ws.append([note])
    for col, width in zip("ABCDE", (20, 10, 22, 90, 50)):
        help_ws.column_dimensions[col].width = width
    for row in help_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    cat_ws = wb.create_sheet("categories")
    _header(cat_ws, ("path", "label", "app_path"))
    for c in cats:
        cat_ws.append([c.path, c.label, c.app_path])
    cat_ws.column_dimensions["A"].width = 44
    cat_ws.column_dimensions["B"].width = 28

    reg_ws = wb.create_sheet("regions")
    _header(
        reg_ws,
        ("region", "name", "currency", "decimals", "tax", "standard rate %", "prices include tax", "you sell here", "your delivery fee"),
    )
    for r in ordered:
        s = mine.get(r.region)
        reg_ws.append(
            [
                r.region, r.name, r.currency, r.exponent, r.tax_name, _pct(r.tax_rate_bp),
                "true" if r.prices_include_tax else "false", "yes" if s else "no",
                _major(s.default_delivery_fee, r.exponent) if s else None,
            ]
        )  # fmt: skip
    reg_ws.column_dimensions["B"].width = 26

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
