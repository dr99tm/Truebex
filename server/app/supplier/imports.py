"""Bulk import in the portal (PF8 Scope 4): upload an XLSX (the template),
a CSV or a JSON feed → a dry run (what would be created, updated, left
unchanged and rejected, with row and column) kept 24 h → Apply, which runs
PF7's importer and records a `feed_runs` row.

The XLSX reader turns the `catalogue` sheet into §6.4 rows before the
importer sees it: spreadsheet numbers are read through `Decimal(str(cell))`
(never a float), so `1299.999` stays three decimals and is refused as
`bad_price`; whole numbers in integer columns lose their `.0`. SKUs, variant
ids and GTINs a spreadsheet turned into numbers (`00123` → `123`, a GTIN in
exponent form) are imported as they read but reported as warnings.
"""

import csv
import io
import re
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..contract_http import ContractError
from ..market import importer
from ..market.common import new_id, not_found, now, rfc3339
from ..market.importer import COLUMNS, REQUIRED_HEADER, FeedInvalid, FeedTooLarge, Row
from ..market.models import FeedRun, Supplier
from ..storage import get_store
from .common import Member, conflict
from .models import SupplierImport

DRY_RUN_TTL = timedelta(hours=24)
FORMATS = ("xlsx", "csv", "json")
# Unpacked size cap of an XLSX (a zip): refuses decompression bombs.
MAX_XLSX_UNPACKED = 400 * 1024 * 1024
INTEGER_COLUMNS = frozenset(
    {"width_mm", "height_mm", "depth_mm", "delivery_days_min", "delivery_days_max", "stock", "lead_time_days"}
)
IDENTITY_COLUMNS = ("sku", "variant_id", "gtin")
_EXPONENT_FORM = re.compile(r"^\d+(\.\d+)?[eE][+-]?\d+$")


# --- Reading ---------------------------------------------------------------------------------


def _xlsx_cell(column: str, value, number: int, warnings: list[dict]) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float, Decimal)):
        d = Decimal(str(value))
        whole = d == d.to_integral_value()
        if column in IDENTITY_COLUMNS:
            warnings.append({"row": number, "column": column, "code": "looks_reformatted", "value": str(value)[:100]})
            return str(int(d)) if whole else format(d, "f")
        if column in INTEGER_COLUMNS and whole:
            return str(int(d))
        return format(d, "f")
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def read_xlsx(data: bytes) -> tuple[list[Row], list[dict]]:
    """Rows of the `catalogue` sheet (or the first sheet), numbered as the
    spreadsheet shows them (the header is row 1)."""
    from openpyxl import load_workbook

    if len(data) > importer.MAX_BYTES:
        raise FeedTooLarge("A file is at most 50 MB; split it by category.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_XLSX_UNPACKED:
                raise FeedTooLarge("This workbook unpacks to more than 400 MB; split it by category.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except FeedTooLarge:
        raise
    except Exception:  # not a zip, not a workbook, or a damaged one
        raise FeedInvalid("The file is not an Excel workbook (.xlsx). Save it as .xlsx or CSV UTF-8.")
    try:
        ws = wb["catalogue"] if "catalogue" in wb.sheetnames else wb.worksheets[0]
        values = ws.iter_rows(values_only=True)
        header = next(values, None)
        if not header:
            raise FeedInvalid("The catalogue sheet is empty: row 1 must name the columns.")
        names = [str(h).strip() if h is not None else "" for h in header]
        while names and not names[-1]:
            names.pop()
        dupes = sorted({n for n in names if n and names.count(n) > 1})
        if dupes:
            raise FeedInvalid(f"These columns appear twice in the header: {', '.join(dupes)}.")
        missing = [c for c in REQUIRED_HEADER if c not in names]
        if missing:
            raise FeedInvalid(f"The header is missing required columns: {', '.join(missing)}.")
        warnings = [{"column": n, "code": "unknown_column"} for n in names if n and n not in COLUMNS]
        rows: list[Row] = []
        for number, cells in enumerate(values, start=2):
            if cells is None or all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
                continue
            if len(rows) >= importer.MAX_ROWS:
                raise FeedTooLarge(f"A file holds at most {importer.MAX_ROWS} rows; split it by category.")
            row = {}
            for k, name in enumerate(names):
                if name in COLUMNS:
                    row[name] = _xlsx_cell(name, cells[k] if k < len(cells) else None, number, warnings)
            rows.append(Row(number, row))
        return rows, warnings
    finally:
        wb.close()


def reformat_warnings(rows: list[Row]) -> list[dict]:
    """A SKU, variant id or GTIN in exponent form (`9.5E+12`) was mangled by a
    spreadsheet program before it was saved as CSV."""
    out = []
    for r in rows:
        for column in IDENTITY_COLUMNS:
            if _EXPONENT_FORM.match(r.get(column)):
                out.append({"row": r.number, "column": column, "code": "looks_reformatted", "value": r.get(column)[:100]})
    return out


def to_csv(rows: list[Row]) -> bytes:
    """The rows as a §6.4 CSV, each on its spreadsheet row (blank rows keep
    the numbering), for the importer's run."""
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(COLUMNS)
    line = 2
    for r in rows:
        while line < r.number:
            w.writerow([""] * len(COLUMNS))
            line += 1
        w.writerow([r.cells.get(c, "") for c in COLUMNS])
        line += 1
    return buf.getvalue().encode("utf-8")


def format_of(filename: str, given: str | None) -> str:
    if given:
        if given not in FORMATS:
            raise ContractError(
                "validation_failed", 422, "format: xlsx, csv or json",
                {"fields": [{"field": "format", "in": "body", "message": "xlsx, csv or json"}]},
            )  # fmt: skip
        return given
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    if ext in FORMATS:
        return ext
    raise ContractError(
        "validation_failed", 422, "format: name the file .xlsx, .csv or .json, or send format",
        {"fields": [{"field": "format", "in": "body", "message": "xlsx, csv or json"}]},
    )  # fmt: skip


def parse_upload(data: bytes, fmt: str) -> tuple[list[Row], list[dict], dict, bytes, str]:
    """(rows, warnings, includes, the bytes a run reads, the run's format)."""
    if fmt == "xlsx":
        rows, warnings = read_xlsx(data)
        includes: dict = {}
        stored, run_fmt = to_csv(rows), "csv"
    else:
        rows, warnings, includes = importer.parse(data, fmt)
        stored, run_fmt = data, fmt
    return rows, warnings + reformat_warnings(rows), includes, stored, run_fmt


# --- Dry run and apply ---------------------------------------------------------------------------


def create(db: Session, m: Member, data: bytes, filename: str, fmt: str, mode: str) -> SupplierImport:
    if mode not in ("upsert", "replace"):
        raise ContractError(
            "validation_failed", 422, "mode: upsert or replace",
            {"fields": [{"field": "mode", "in": "body", "message": "upsert or replace"}]},
        )  # fmt: skip
    try:
        rows, warnings, includes, stored, run_fmt = parse_upload(data, fmt)
    except FeedTooLarge as exc:
        raise ContractError("too_large", 413, str(exc))
    except FeedInvalid as exc:
        raise ContractError("feed_invalid", 422, str(exc))
    supplier_id = m.supplier.supplier_id
    preview = importer.dry_run_rows(db, m.supplier, rows, mode, warnings, includes)
    import_id = new_id()
    key = f"market/suppliers/{supplier_id}/imports/{import_id}/source.{run_fmt}"
    get_store().put(key, stored, content_type="text/csv" if run_fmt == "csv" else "application/json")
    imp = SupplierImport(
        import_id=import_id,
        supplier_id=supplier_id,
        filename=(filename or "")[:200],
        source_key=key,
        format=fmt,
        mode=mode,
        dry_run=preview,
        created_by=m.user.id,
        created_at=now(),
    )
    db.add(imp)
    db.commit()
    return imp


def get(db: Session, supplier: Supplier, import_id: str) -> SupplierImport:
    imp = db.get(SupplierImport, import_id) if isinstance(import_id, str) else None
    if imp is None or imp.supplier_id != supplier.supplier_id:
        raise not_found("That import")
    return imp


def expired(imp: SupplierImport) -> bool:
    from ..market.common import aware

    return imp.feed_id is None and aware(imp.created_at) < now() - DRY_RUN_TTL


def apply(db: Session, m: Member, imp: SupplierImport) -> FeedRun:
    if imp.feed_id:
        raise conflict("This upload was applied already.", "already_applied")
    if expired(imp):
        raise ContractError("expired", 410, "This dry run is more than 24 hours old. Upload the file again.")
    with get_store().open(imp.source_key) as fh:
        data = fh.read()
    run_fmt = "json" if imp.format == "json" else "csv"
    try:
        run = importer.queue(db, m.supplier, data, run_fmt, imp.mode, source="portal", by_user_id=m.user.id)
    except (FeedInvalid, FeedTooLarge) as exc:  # validated at upload; kept for safety
        raise ContractError("feed_invalid", 422, str(exc))
    imp = db.get(SupplierImport, imp.import_id)
    imp.feed_id = run.feed_id
    db.add(imp)
    db.commit()
    return run


def import_json(db: Session, imp: SupplierImport) -> dict:
    run = db.get(FeedRun, imp.feed_id) if imp.feed_id else None
    return {
        "import_id": imp.import_id,
        "filename": imp.filename,
        "format": imp.format,
        "mode": imp.mode,
        "dry_run": imp.dry_run or {},
        "feed_id": imp.feed_id,
        "run": importer.report_json(run) if run else None,
        "created_at": rfc3339(imp.created_at),
        "expires_at": rfc3339(imp.created_at + DRY_RUN_TTL) if not imp.feed_id else None,
        "expired": expired(imp),
    }


def history(db: Session, supplier: Supplier) -> dict:
    imports = db.scalars(
        select(SupplierImport)
        .where(SupplierImport.supplier_id == supplier.supplier_id)
        .order_by(SupplierImport.created_at.desc())
        .limit(50)
    )
    runs = db.scalars(
        select(FeedRun).where(FeedRun.supplier_id == supplier.supplier_id).order_by(FeedRun.created_at.desc()).limit(50)
    )
    return {
        "imports": [import_json(db, i) for i in imports],
        "runs": [importer.report_json(r) for r in runs],
    }


def purge(db: Session, at: datetime) -> int:
    """Drop dry runs never applied, 24 h after the upload (and their file)."""
    from ..market.common import aware

    store = get_store()
    n = 0
    for imp in list(db.scalars(select(SupplierImport).where(SupplierImport.feed_id.is_(None)))):
        if aware(imp.created_at) < at - DRY_RUN_TTL:
            try:
                store.delete(imp.source_key)
            except Exception:  # noqa: BLE001 - a missing file is fine
                pass
            db.delete(imp)
            n += 1
    db.commit()
    return n
